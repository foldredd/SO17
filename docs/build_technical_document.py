from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SO17_Technical_Documentation.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.line_spacing = 1.14
for style_name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
    style = styles[style_name]
    style.font.name = 'Aptos Display' if style_name == 'Title' else 'Aptos'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), style.font.name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(16 if style_name != 'Title' else 0)
    style.paragraph_format.space_after = Pt(7)

title_style_props = styles['Title']._element.find(qn('w:pPr'))
if title_style_props is not None:
    for title_style_border in title_style_props.findall(qn('w:pBdr')):
        title_style_props.remove(title_style_border)

def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement('w:shd')
    fill.set(qn('w:fill'), color)
    props.append(fill)

def borders(table):
    tbl_pr = table._tbl.tblPr
    border = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')
        el.set(qn('w:color'), 'D9D9D9')
        border.append(el)
    tbl_pr.append(border)

def table(headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    borders(tbl)
    for idx, heading in enumerate(headers):
        cell = tbl.rows[0].cells[idx]
        cell.text = heading
        shade(cell, '1F4E78')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
    header_props = tbl.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement('w:tblHeader')
    repeat_header.set(qn('w:val'), 'true')
    header_props.append(repeat_header)
    for row_index, values in enumerate(rows):
        cells = tbl.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[idx], 'F3F7FA')
            for paragraph in cells[idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
        if widths:
            for idx, width in enumerate(widths):
                cells[idx].width = Inches(width)
    doc.add_paragraph()
    return tbl

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    run.font.size = Pt(9.5)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('Техническая документация SO17')
title_props = title._p.get_or_add_pPr()
title_border = title_props.find(qn('w:pBdr'))
if title_border is not None:
    title_props.remove(title_border)
p = doc.add_paragraph()
p.add_run('Версия разработки  ').bold = True
p.add_run('14 сентября 2026')
p = doc.add_paragraph()
p.add_run('Назначение. ').bold = True
p.add_run('Документ объясняет, как устроен обновлённый веб проект SO17, как его запускать и проверять. Он рассчитан на владельца проекта и начинающего разработчика.')

doc.add_heading('Результат работы', level=1)
doc.add_paragraph('Проект разделён на независимые React frontend и ASP.NET Core backend. Данные хранятся в SQL Server Express через Entity Framework Core. Запуск автоматизирован: один сценарий подготавливает локальную конфигурацию, другой применяет миграции и запускает оба приложения.')
table(['Часть', 'Папка', 'Ответственность'], [
    ('Frontend', 'frontend', 'Интерфейс React, маршруты, формы, таблицы и клиент API.'),
    ('API', 'backend/SO17.Api', 'HTTP endpoints, вход, защита запросов, обработка ошибок и конфигурация.'),
    ('Application', 'backend/SO17.Application', 'Сценарии работы приложения, валидация и контракты.'),
    ('Domain', 'backend/SO17.Domain', 'Сущности и правила предметной области: заказы, остатки и статусы.'),
    ('Infrastructure', 'backend/SO17.Infrastructure', 'Entity Framework, SQL Server, Identity, репозитории и миграции.'),
    ('Tests', 'tests/SO17.Tests', 'Интеграционные тесты API с отдельной временной базой SQL Server.')
], [1.0, 1.75, 3.7])

doc.add_page_break()
doc.add_heading('Быстрый запуск', level=1)
doc.add_paragraph('SQL Server Express должен быть запущен. В обычном случае служба MSSQL$SQLEXPRESS стартует вместе с Windows. В PowerShell из корневой папки проекта выполните следующие команды:')
code('powershell -NoProfile -ExecutionPolicy Bypass -File .\\scripts\\Setup-Dev.ps1')
code('powershell -NoProfile -ExecutionPolicy Bypass -File .\\scripts\\Start-Dev.ps1')
doc.add_paragraph('После сообщения READY откройте http://localhost:5173. При первой регистрации понадобится код приглашения из файла backend\\SO17.Api\\appsettings.Local.json. Этот локальный файл создаётся автоматически и не отслеживается Git.')
table(['Проверка', 'Адрес или команда', 'Ожидаемый результат'], [
    ('Интерфейс', 'http://localhost:5173', 'Открывается экран входа или регистрации.'),
    ('API', 'http://localhost:5080/openapi/v1.json', 'Возвращается описание HTTP API.'),
    ('База', 'http://localhost:5080/health/ready', 'Текст Healthy.'),
    ('Остановка', '.\\scripts\\Stop-Dev.ps1', 'Оба фоновых процесса завершаются.')
], [1.1, 2.75, 2.6])

doc.add_heading('Модель данных и правила', level=1)
doc.add_paragraph('В SQL Server используется база SO17_Refactored. Миграции Entity Framework создают таблицы пользователей Identity, клиентов, товаров, заказов и строк заказов. Никаких ручных SQL скриптов для создания таблиц выполнять не нужно.')
table(['Сущность', 'Ключевые данные', 'Правило'], [
    ('Customer', 'Имя, фамилия, телефон', 'Телефон уникален.'),
    ('Item', 'Название, SKU, количество, цена', 'SKU уникален; количество не может стать отрицательным.'),
    ('Order', 'Клиент, статус, сумма', 'Разрешён переход Pending → Confirmed или Cancelled.'),
    ('OrderLine', 'Снимок названия, SKU, цены, количества', 'Цена и название фиксируются на момент создания заказа.')
], [1.05, 2.65, 2.75])
doc.add_paragraph('Для сущностей используется поле Version. Оно защищает от потери данных при одновременном редактировании: пользователь получит понятную ошибку конфликта вместо тихой перезаписи чужих изменений. Создание заказа и списание остатков выполняются в одной транзакции. При отмене товара возвращаются один раз.')

doc.add_heading('Безопасность', level=1)
table(['Риск в исходном подходе', 'Исправление'], [
    ('Секретный ключ мог находиться в исходном коде или общем конфигурационном файле.', 'Код приглашения создаётся в appsettings.Local.json. Файл внесён в .gitignore.'),
    ('Пароли можно было хранить или сравнивать небезопасно.', 'Используется ASP.NET Core Identity и безопасное хеширование паролей.'),
    ('Сессия могла быть доступна JavaScript или уйти на другой сайт.', 'Cookie имеет HttpOnly, SameSite Strict и Secure в production.'),
    ('Изменяющие запросы могли выполняться с чужого сайта.', 'Для POST, PUT, DELETE требуется CSRF заголовок и cookie.'),
    ('Открытые endpoints и CORS расширяли доступ.', 'По умолчанию требуется вход. В production разрешаются только явно указанные origins.'),
    ('Ошибки БД и исключения могли раскрываться пользователю.', 'Problem Details возвращает безопасный ответ с trace id; подробность записывается в журнал.'),
    ('Подбор паролей мог быть неограниченным.', 'На входе действует rate limit; у Identity есть блокировка после неудачных попыток.')
], [2.85, 3.6])

doc.add_heading('Применённые принципы', level=1)
doc.add_paragraph('SOLID реализован через разделение обязанностей. Контроллеры принимают HTTP запрос и не содержат бизнес логику. Сервисы Application реализуют сценарии, репозитории Infrastructure изолируют Entity Framework, а Domain не зависит от базы данных или веб фреймворка. Зависимости передаются через интерфейсы и контейнер ASP.NET Core.')
bullet('Single Responsibility: API, бизнес правила, сохранение данных и интерфейс размещены в разных проектах и файлах.')
bullet('Open Closed: новые реализации хранилища или внешние сервисы можно добавлять через интерфейсы без переписывания контроллеров.')
bullet('Dependency Inversion: сценарии зависят от абстракций репозиториев и Unit of Work, а не от AppDbContext.')

doc.add_heading('API и пользовательские возможности', level=1)
doc.add_paragraph('Frontend предоставляет страницы входа, панели управления, клиентов, товаров, заказов и профиля. API поддерживает регистрацию по коду приглашения, вход и выход, смену пароля, CRUD операции для клиентов и товаров, создание заказа, смену статуса, поиск и постраничную выдачу.')
table(['Группа endpoint', 'Назначение'], [
    ('/api/account', 'Регистрация, вход, выход, профиль и пароль.'),
    ('/api/customers', 'Список, поиск, создание, изменение и удаление клиентов.'),
    ('/api/items', 'Работа с каталогом и остатками товаров.'),
    ('/api/orders', 'Создание, список, детали, подтверждение и отмена заказов.'),
    ('/api/dashboard', 'Сводные показатели для панели управления.')
], [1.65, 4.85])

doc.add_heading('Проверки качества', level=1)
doc.add_paragraph('Проверки запускаются командой ниже. Она не меняет рабочую базу данных: интеграционные тесты создают и удаляют собственную временную базу с уникальным именем.')
code('powershell -NoProfile -ExecutionPolicy Bypass -File .\\scripts\\Test.ps1')
table(['Проверка', 'Результат'], [
    ('Сборка .NET', '0 предупреждений и 0 ошибок.'),
    ('Интеграционные тесты', '15 из 15 пройдены: доступ, CSRF, валидация, CRUD, уникальность, остатки, статусы, конфликт параллельных изменений, сессия и вход.'),
    ('Frontend lint', 'Проверка ESLint проходит.'),
    ('Frontend production build', 'Vite успешно формирует dist.'),
    ('Уязвимые .NET пакеты', 'dotnet list package --vulnerable не обнаружил уязвимых пакетов.')
], [2.25, 4.25])

doc.add_heading('Диагностика', level=1)
doc.add_paragraph('Если SQL Server остановлен, откройте PowerShell от имени администратора и выполните Start-Service \'MSSQL$SQLEXPRESS\'. Состояние службы проверяется командой Get-Service \'MSSQL$SQLEXPRESS\'. Если порт занят, сначала выполните Stop-Dev.ps1. Журналы API и frontend находятся в папке .local; она локальная и не передаётся в Git.')
doc.add_paragraph('Для production необходимо задать собственные защищённые значения ConnectionStrings, Registration и DataProtection через секретное хранилище или переменные окружения, указать разрешённые Cors origins и использовать HTTPS. Локальный код приглашения нельзя публиковать.')

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('SO17  Техническая документация')
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)

doc.save(OUT)
print(OUT)
